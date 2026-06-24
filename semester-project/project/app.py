from flask import Flask, render_template, request, jsonify, send_file
from flask_sqlalchemy import SQLAlchemy
from datetime import datetime
import os
from io import BytesIO
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

app = Flask(__name__)

basedir = os.path.abspath(os.path.dirname(__file__))

app.config["SQLALCHEMY_DATABASE_URI"] = "sqlite:///" + os.path.join(basedir, "warehouse.db")
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False

db = SQLAlchemy(app)


class Category(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(120), unique=True, nullable=False)
    products = db.relationship("Product", backref="category", lazy=True)


class Product(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(200), nullable=False)
    category_id = db.Column(db.Integer, db.ForeignKey("category.id"), nullable=False)
    unit = db.Column(db.String(50), default="шт")
    price = db.Column(db.Float, default=0)

    batches = db.relationship("Batch", backref="product", lazy=True)

    def total_quantity(self):
        return sum(b.quantity for b in self.batches)


class Storage(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    building = db.Column(db.String(50), nullable=False)
    department = db.Column(db.String(50), nullable=False)
    shelf = db.Column(db.String(50), nullable=False)
    capacity = db.Column(db.Integer, default=1000)
    current_quantity = db.Column(db.Integer, default=0)

    batches = db.relationship("Batch", backref="storage", lazy=True)

    def address(self):
        return f"Корпус {self.building}, відділення {self.department}, полиця {self.shelf}"


class Batch(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    product_id = db.Column(db.Integer, db.ForeignKey("product.id"), nullable=False)
    storage_id = db.Column(db.Integer, db.ForeignKey("storage.id"), nullable=False)
    quantity = db.Column(db.Integer, nullable=False)
    batch_number = db.Column(db.String(100), unique=True, nullable=False)
    supplier = db.Column(db.String(200), default="")
    receipt_date = db.Column(db.DateTime, default=datetime.utcnow)

    invoices = db.relationship("Invoice", backref="batch", lazy=True)


class Invoice(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    batch_id = db.Column(db.Integer, db.ForeignKey("batch.id"), nullable=False)
    invoice_number = db.Column(db.String(100), unique=True, nullable=False)
    quantity_released = db.Column(db.Integer, nullable=False)
    recipient = db.Column(db.String(200), default="")
    release_date = db.Column(db.DateTime, default=datetime.utcnow)


def add_category_if_missing(name):
    category = Category.query.filter_by(name=name).first()

    if not category:
        category = Category(name=name)
        db.session.add(category)
        db.session.commit()

    return category


def add_storage_if_missing(building, department, shelf, capacity):
    storage = Storage.query.filter_by(
        building=building,
        department=department,
        shelf=shelf
    ).first()

    if not storage:
        storage = Storage(
            building=building,
            department=department,
            shelf=shelf,
            capacity=capacity
        )
        db.session.add(storage)
        db.session.commit()

    return storage


def add_product_if_missing(name, category_name, unit, price):
    product = Product.query.filter_by(name=name).first()

    if product:
        return product

    category = Category.query.filter_by(name=category_name).first()

    if not category:
        category = add_category_if_missing(category_name)

    product = Product(
        name=name,
        category_id=category.id,
        unit=unit,
        price=price
    )

    db.session.add(product)
    db.session.commit()

    return product


def add_batch_if_missing(product_name, storage, quantity, batch_number, supplier):
    batch = Batch.query.filter_by(batch_number=batch_number).first()

    if batch:
        return batch

    product = Product.query.filter_by(name=product_name).first()

    if not product:
        return None

    if storage.current_quantity + quantity <= storage.capacity:
        batch = Batch(
            product_id=product.id,
            storage_id=storage.id,
            quantity=quantity,
            batch_number=batch_number,
            supplier=supplier
        )

        storage.current_quantity += quantity

        db.session.add(batch)
        db.session.commit()

        return batch

    return None


with app.app_context():
    db.create_all()

    start_categories = [
        "Техніка",
        "Комп'ютерна техніка",
        "Смартфони",
        "Ноутбуки",
        "Побутова техніка",
        "Ліки",
        "Медичні товари",
        "Одяг",
        "Взуття",
        "Косметика",
        "Харчові продукти",
        "Напої",
        "Канцелярія",
        "Книги",
        "Іграшки",
        "Меблі",
        "Будівельні матеріали",
        "Автотовари",
        "Спортивні товари",
        "Електроінструменти",
        "Сад та город",
        "Товари для тварин"
    ]

    for name in start_categories:
        add_category_if_missing(name)

    storage_a1 = add_storage_if_missing("А", "1", "1", 500)
    storage_a2 = add_storage_if_missing("А", "1", "2", 500)
    storage_a3 = add_storage_if_missing("А", "2", "1", 700)
    storage_b1 = add_storage_if_missing("Б", "1", "1", 1000)
    storage_b2 = add_storage_if_missing("Б", "2", "3", 800)
    storage_v1 = add_storage_if_missing("В", "1", "5", 600)
    storage_v2 = add_storage_if_missing("В", "3", "2", 900)

    add_product_if_missing("Ноутбук ASUS TUF Gaming", "Ноутбуки", "шт", 42000)
    add_product_if_missing("Монітор Samsung 27", "Техніка", "шт", 8500)
    add_product_if_missing("iPhone 15", "Смартфони", "шт", 39999)
    add_product_if_missing("Samsung Galaxy S24", "Смартфони", "шт", 35999)
    add_product_if_missing("Парацетамол", "Ліки", "уп", 65)
    add_product_if_missing("Амоксицилін", "Ліки", "уп", 145)
    add_product_if_missing("Кава Jacobs", "Харчові продукти", "шт", 220)
    add_product_if_missing("Чай Greenfield", "Напої", "шт", 95)
    add_product_if_missing("Зошит 48 арк.", "Канцелярія", "шт", 35)
    add_product_if_missing("Ручка Pilot", "Канцелярія", "шт", 25)
    add_product_if_missing("Худі Oversize", "Одяг", "шт", 1200)
    add_product_if_missing("Джинси Classic", "Одяг", "шт", 1800)
    add_product_if_missing("Шампунь Elseve", "Косметика", "шт", 160)
    add_product_if_missing("Корм для котів Whiskas", "Товари для тварин", "шт", 180)

    add_batch_if_missing("Ноутбук ASUS TUF Gaming", storage_a1, 8, "BATCH-001", "ASUS Ukraine")
    add_batch_if_missing("Монітор Samsung 27", storage_a2, 15, "BATCH-002", "Samsung Store")
    add_batch_if_missing("iPhone 15", storage_a3, 10, "BATCH-003", "Apple Partner")
    add_batch_if_missing("Samsung Galaxy S24", storage_a3, 12, "BATCH-004", "Samsung Store")
    add_batch_if_missing("Парацетамол", storage_b1, 120, "BATCH-005", "Аптека Склад")
    add_batch_if_missing("Амоксицилін", storage_b1, 80, "BATCH-006", "МедПостач")
    add_batch_if_missing("Кава Jacobs", storage_b2, 60, "BATCH-007", "Food Market")
    add_batch_if_missing("Чай Greenfield", storage_b2, 75, "BATCH-008", "Food Market")
    add_batch_if_missing("Зошит 48 арк.", storage_v1, 200, "BATCH-009", "КанцОпт")
    add_batch_if_missing("Ручка Pilot", storage_v1, 300, "BATCH-010", "КанцОпт")
    add_batch_if_missing("Худі Oversize", storage_v2, 25, "BATCH-011", "Fashion Shop")
    add_batch_if_missing("Джинси Classic", storage_v2, 20, "BATCH-012", "Fashion Shop")
    add_batch_if_missing("Шампунь Elseve", storage_b2, 40, "BATCH-013", "Beauty Market")
    add_batch_if_missing("Корм для котів Whiskas", storage_b2, 55, "BATCH-014", "Zoo Market")


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/stats")
def api_stats():
    return jsonify({
        "products": Product.query.count(),
        "categories": Category.query.count(),
        "batches": Batch.query.count(),
        "storage": Storage.query.count(),
        "invoices": Invoice.query.count()
    })


@app.route("/api/categories", methods=["GET", "POST"])
def api_categories():
    if request.method == "POST":
        data = request.get_json()

        name = data.get("name", "").strip()

        if not name:
            return jsonify({"error": "Введіть назву категорії"}), 400

        if Category.query.filter_by(name=name).first():
            return jsonify({"error": "Така категорія вже існує"}), 400

        category = Category(name=name)

        db.session.add(category)
        db.session.commit()

        return jsonify({
            "id": category.id,
            "name": category.name
        }), 201

    categories = Category.query.order_by(Category.name).all()

    return jsonify([
        {
            "id": c.id,
            "name": c.name
        }
        for c in categories
    ])


@app.route("/api/products", methods=["GET", "POST"])
def api_products():
    if request.method == "POST":
        data = request.get_json()

        name = data.get("name", "").strip()
        category_id = data.get("category_id")
        unit = data.get("unit", "шт").strip()
        price = float(data.get("price", 0))

        if not name:
            return jsonify({"error": "Введіть назву товару"}), 400

        category = Category.query.get(category_id)

        if not category:
            return jsonify({"error": "Категорію не знайдено"}), 400

        product = Product(
            name=name,
            category_id=category_id,
            unit=unit,
            price=price
        )

        db.session.add(product)
        db.session.commit()

        return jsonify({
            "id": product.id,
            "name": product.name
        }), 201

    search = request.args.get("search", "").strip()
    category_id = request.args.get("category_id", "").strip()

    query = Product.query

    if search:
        query = query.filter(Product.name.ilike(f"%{search}%"))

    if category_id:
        query = query.filter(Product.category_id == int(category_id))

    products = query.order_by(Product.id.desc()).all()

    return jsonify([
        {
            "id": p.id,
            "name": p.name,
            "category": p.category.name,
            "category_id": p.category_id,
            "unit": p.unit,
            "price": p.price,
            "quantity": p.total_quantity()
        }
        for p in products
    ])


@app.route("/api/storage")
def api_storage():
    storages = Storage.query.order_by(
        Storage.building,
        Storage.department,
        Storage.shelf
    ).all()

    return jsonify([
        {
            "id": s.id,
            "address": s.address(),
            "capacity": s.capacity,
            "used": s.current_quantity,
            "available": s.capacity - s.current_quantity
        }
        for s in storages
    ])


@app.route("/api/batches", methods=["GET", "POST"])
def api_batches():
    if request.method == "POST":
        data = request.get_json()

        product_id = data.get("product_id")
        storage_id = data.get("storage_id")
        quantity = int(data.get("quantity", 0))
        batch_number = data.get("batch_number", "").strip()
        supplier = data.get("supplier", "").strip()

        product = Product.query.get(product_id)
        storage = Storage.query.get(storage_id)

        if not product:
            return jsonify({"error": "Товар не знайдено"}), 400

        if not storage:
            return jsonify({"error": "Місце зберігання не знайдено"}), 400

        if quantity <= 0:
            return jsonify({"error": "Кількість має бути більше нуля"}), 400

        if not batch_number:
            return jsonify({"error": "Введіть номер партії"}), 400

        if Batch.query.filter_by(batch_number=batch_number).first():
            return jsonify({"error": "Партія з таким номером вже існує"}), 400

        if storage.current_quantity + quantity > storage.capacity:
            return jsonify({"error": "Недостатньо місця на складі"}), 400

        batch = Batch(
            product_id=product_id,
            storage_id=storage_id,
            quantity=quantity,
            batch_number=batch_number,
            supplier=supplier
        )

        storage.current_quantity += quantity

        db.session.add(batch)
        db.session.commit()

        return jsonify({
            "id": batch.id,
            "batch_number": batch.batch_number
        }), 201

    product_id = request.args.get("product_id", "").strip()

    query = Batch.query

    if product_id:
        query = query.filter(Batch.product_id == int(product_id))

    batches = query.order_by(Batch.id.desc()).all()

    return jsonify([
        {
            "id": b.id,
            "product_id": b.product_id,
            "product": b.product.name,
            "category": b.product.category.name,
            "batch_number": b.batch_number,
            "quantity": b.quantity,
            "unit": b.product.unit,
            "storage": b.storage.address(),
            "supplier": b.supplier,
            "date": b.receipt_date.strftime("%d.%m.%Y")
        }
        for b in batches
    ])


@app.route("/api/release", methods=["POST"])
def api_release():
    data = request.get_json()

    batch_id = data.get("batch_id")
    quantity = int(data.get("quantity", 0))
    recipient = data.get("recipient", "").strip()

    batch = Batch.query.get(batch_id)

    if not batch:
        return jsonify({"error": "Партію не знайдено"}), 400

    if quantity <= 0:
        return jsonify({"error": "Кількість має бути більше нуля"}), 400

    if batch.quantity < quantity:
        return jsonify({"error": "Недостатньо товару в цій партії"}), 400

    if not recipient:
        return jsonify({"error": "Введіть одержувача"}), 400

    invoice_number = f"INV-{datetime.now().strftime('%Y%m%d%H%M%S')}"

    invoice = Invoice(
        batch_id=batch.id,
        invoice_number=invoice_number,
        quantity_released=quantity,
        recipient=recipient
    )

    batch.quantity -= quantity
    batch.storage.current_quantity -= quantity

    db.session.add(invoice)
    db.session.commit()

    return jsonify({
        "invoice_id": invoice.id,
        "invoice_number": invoice.invoice_number
    }), 201


@app.route("/api/invoices")
def api_invoices():
    invoices = Invoice.query.order_by(Invoice.release_date.desc()).all()

    return jsonify([
        {
            "id": i.id,
            "invoice_number": i.invoice_number,
            "product": i.batch.product.name,
            "quantity": i.quantity_released,
            "unit": i.batch.product.unit,
            "recipient": i.recipient,
            "date": i.release_date.strftime("%d.%m.%Y %H:%M")
        }
        for i in invoices
    ])


@app.route("/api/invoice/<int:invoice_id>/download")
def download_invoice(invoice_id):
    invoice = Invoice.query.get(invoice_id)

    if not invoice:
        return jsonify({"error": "Накладну не знайдено"}), 404

    batch = invoice.batch
    product = batch.product

    doc = Document()

    title = doc.add_heading("НАКЛАДНА НА ВІДПУСК ТОВАРУ", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(f"Номер накладної: {invoice.invoice_number}")
    doc.add_paragraph(f"Дата: {invoice.release_date.strftime('%d.%m.%Y %H:%M')}")
    doc.add_paragraph("")

    doc.add_paragraph("Інформація про товар:")
    doc.add_paragraph(f"Назва товару: {product.name}")
    doc.add_paragraph(f"Категорія: {product.category.name}")
    doc.add_paragraph(f"Кількість: {invoice.quantity_released} {product.unit}")
    doc.add_paragraph(f"Ціна за одиницю: {product.price} грн")
    doc.add_paragraph(f"Сума: {invoice.quantity_released * product.price} грн")
    doc.add_paragraph("")

    doc.add_paragraph("Інформація про партію:")
    doc.add_paragraph(f"Номер партії: {batch.batch_number}")
    doc.add_paragraph(f"Місце зберігання: {batch.storage.address()}")
    doc.add_paragraph(f"Постачальник: {batch.supplier}")
    doc.add_paragraph("")

    doc.add_paragraph(f"Одержувач: {invoice.recipient}")
    doc.add_paragraph("")
    doc.add_paragraph("Відпустив: ____________________")
    doc.add_paragraph("Отримав: ____________________")
    doc.add_paragraph("")
    doc.add_paragraph("Роботу виконала: Рибакова Альона Олексіївна")
    doc.add_paragraph("Статистика, 1 курс, 2 семестр")

    output = BytesIO()
    doc.save(output)
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name=f"{invoice.invoice_number}.docx",
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )


if __name__ == "__main__":
    app.run(debug=True, host="127.0.0.1", port=5000)